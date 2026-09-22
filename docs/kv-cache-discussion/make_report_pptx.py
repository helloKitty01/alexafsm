#!/usr/bin/env python3
"""把《kv cache方案（汇报版）》HTML 幻灯片导出为真实 PPT(.pptx,16:9,3 页)。

数字不手填:脚本从 HTML 里的 `BASELINE` 数据块(单一事实来源)用 node 算出全部结果再排版,
HTML 改基线,重跑本脚本即同步。

用法:
    python3 make_report_pptx.py [源 HTML] [输出 .pptx]
默认:源 = 目录下版本号最高的 `kv cache方案（汇报版）.v*.slides.html`,输出 = `kv cache方案（汇报版）.v2.pptx`

依赖:python-pptx、node(≥14)。
"""
import glob
import json
import os
import re
import subprocess
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

HERE = os.path.dirname(os.path.abspath(__file__))

# ---------------- 色板(与 HTML :root 一致) ----------------
TEXT = RGBColor(0x1F, 0x2A, 0x3D)
MUTED = RGBColor(0x5C, 0x6B, 0x84)
ACCENT = RGBColor(0x25, 0x63, 0xEB)
GOOD = RGBColor(0x05, 0x96, 0x69)
WARN = RGBColor(0xD9, 0x77, 0x06)
BAD = RGBColor(0xDC, 0x26, 0x26)
PROD = RGBColor(0x3B, 0x5B, 0xDB)
SESS = RGBColor(0x7C, 0x3A, 0xED)
QLV = RGBColor(0xD9, 0x77, 0x06)
LOOP = RGBColor(0x08, 0x91, 0xB2)
CARD = RGBColor(0xF7, 0xF9, 0xFC)
CARD_BD = RGBColor(0xDF, 0xE5, 0xEE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xFF, 0xFF, 0xFF)
HIT = RGBColor(0x4E, 0xB3, 0x8F)      # 已缓存
RW = RGBColor(0xE0, 0x3C, 0x3C)       # 原位重写
DRAG = RGBColor(0xF5, 0xA5, 0x2B)     # 被连带 · 白算
NEW = RGBColor(0x3F, 0x76, 0xE8)      # 新增 · 必算
FIRST = RGBColor(0x9D, 0xBB, 0xF3)    # 会话首调载入
NA = RGBColor(0xF1, 0xF4, 0xF9)
MV_BG = RGBColor(0xE8, 0xEE, 0xFB)    # 位置变了的行
SUM_BG = RGBColor(0xEA, 0xF0, 0xFC)
BANNER_BG = RGBColor(0xEC, 0xF2, 0xFC)
BANNER_BD = RGBColor(0xA9, 0xC1, 0xF3)
GRID = RGBColor(0xE3, 0xE8, 0xF0)

FONT = "Microsoft YaHei"
MONO = "Consolas"

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)


# ---------------- 数据:从 HTML 的 BASELINE 用 node 算 ----------------
JS_TAIL = r"""
(function (B) {
  function r1(x) { return Math.round(x * 10) / 10; }
  var out = {
    base: { tools: B.tools, sp: B.sp, sess: B.sess, hist: B.hist, dyn: B.dyn, u: B.u, call: B.call, web: B.web, skill: B.skill, tool: B.tool, ans: B.ans, mix: B.mix },
    alpha: B.alpha, Tn: B.Tn, T1: B.T1, fixed2: B.fixed2, C2base: B.C2base, qFirstNew: B.qFirstNew, firstNew: B.firstNew, firstExtra: B.firstExtra, loopW: B.loopW,
    msPerK: B.msPerK, preheatMs: B.preheatMs, trimAt: B.trimAt, costFactor: B.costFactor,
    sessAvg: B.sessAvg, dayAvgRounds: B.dayAvgRounds, sessionN: B.sessionN,
    rounds: B.rounds, p1Groups: B.p1Groups,
    agg: { saveRate: B.agg.saveRate, costSave: B.agg.costSave, curHitAll: B.agg.curHitAll, tgtHitAll: B.agg.tgtHitAll, curTot: B.agg.curTot, tgtTot: B.agg.tgtTot,
           share1: B.agg.share1, longShare: B.agg.longShare, longShareCost: B.agg.longShareCost, longShareSave: B.agg.longShareSave,
           groups: B.agg.groups },
    perRound: [1, 2, 3, 4, 5, 6].map(function (n) { return { n: n, C: B.C(n), T: B.T(n), eta: B.eta(n), hitCur: B.hit(n, 'cur'), hitTgt: B.hit(n, 'tgt') }; }),
    sessN: [1, 2, 3, 4, 5, 6, 8, 10].map(function (n) { var o = B.sessN(n); o.n = n; return o; }),
    rounds7: [1, 2, 3, 4, 5, 6, 10].map(function (n) { return { n: n, C: B.C(n), T: B.T(n), eta: B.eta(n), hitCur: B.hit(n, 'cur'), hitTgt: B.hit(n, 'tgt'), fCur: B.firstCur(n), fTgt: B.firstTgt(n), fCurS: B.sec(B.firstCur(n)), fTgtS: B.sec(B.firstTgt(n)) }; }),
    qNew: B.qNew, loopNew: B.loopNew,
    demo: B.demo,
    hit: { h1: B.hit(1, 'tgt'), h6: B.hit(6, 'tgt'), h12: B.hit(12, 'tgt') }
  };
  console.log(JSON.stringify(out));
})(BASELINE);
"""


def load_data(html_path):
    html = open(html_path, encoding="utf-8").read()
    m = re.search(r"<script>(.*?)/\* =+ 渲染 =+ \*/", html, re.S)
    if not m:
        raise SystemExit("找不到 BASELINE 数据块")
    js = m.group(1) + JS_TAIL
    res = subprocess.run(["node", "-e", js], capture_output=True, text=True)
    if res.returncode != 0:
        raise SystemExit("node 执行失败:\n" + res.stderr)
    return json.loads(res.stdout.strip().splitlines()[-1])


# ---------------- 绘图小工具 ----------------
import math


def K1(x):
    return f"{math.floor(x * 10 + 0.5) / 10:g}K"   # 四舍五入(避开 Python 银行家舍入,与 HTML 的 Math.round 一致)


def K0(x):
    return f"{math.floor(x + 0.5)}K"


def pct(x):
    return f"{math.floor(x * 100 + 0.5)}%"


def pct1(x):
    return f"{x * 100:.1f}%"


def rgb_hex(c):
    return RGBColor.from_string(c.lstrip("#"))


def rect(slide, x, y, w, h, fill=CARD, line=CARD_BD, radius=0.06, lw=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    s.text_frame.margin_left = s.text_frame.margin_right = Inches(0.04)
    s.text_frame.margin_top = s.text_frame.margin_bottom = Inches(0.02)
    return s


def text(slide, x, y, w, h, runs, size=9, color=TEXT, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         mono=False, wrap=True, margin=0.03, line_spacing=1.0, shape=None):
    """runs: str 或 [(text, {size,color,bold,mono}), ...] 或 段落列表 [[runs...], [runs...]]"""
    box = shape if shape is not None else slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.015)
    paras = runs
    if isinstance(runs, str):
        paras = [[(runs, {})]]
    elif runs and isinstance(runs[0], tuple):
        paras = [runs]
    for pi, para in enumerate(paras):
        p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        if isinstance(para, str):
            para = [(para, {})]
        for t, st in para:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = MONO if st.get("mono", mono) else FONT
            f.size = Pt(st.get("size", size))
            f.bold = st.get("bold", bold)
            f.color.rgb = st.get("color", color)
    return box


def card(slide, x, y, w, h, title=None, sub=None, title_size=10.5):
    rect(slide, x, y, w, h)
    if title:
        runs = [(title, {"bold": True, "size": title_size, "color": TEXT})]
        if sub:
            runs.append(("  " + sub, {"size": title_size - 3, "color": MUTED}))
        text(slide, x + Inches(0.1), y + Inches(0.05), w - Inches(0.2), Inches(0.3), runs, anchor=MSO_ANCHOR.MIDDLE)
    return y + Inches(0.38 if title else 0.08)


def chip(slide, x, y, w, h, label, fill, color=WHITE, size=6.5, bold=True):
    s = rect(slide, x, y, w, h, fill=fill, line=None, radius=0.25)
    text(slide, 0, 0, 0, 0, [(label, {"size": size, "bold": bold, "color": color})], align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, margin=0.0, shape=s)
    return s


def table(slide, x, y, w, col_w, rows, header_size=7.5, body_size=7.5, row_h=0.22, hdr_h=None, align=None,
          row_fill=None, cell_style=None, hdr_color=ACCENT, bd=CARD_BD):
    """rows[0] = 表头;每格为 str 或 (str, style-dict);col_w 为比例列表;align 每列 'l'/'r'/'c'。"""
    n, m = len(rows), len(rows[0])
    total = sum(col_w)
    hdr_h = hdr_h or row_h
    h = Inches(hdr_h) + Inches(row_h) * (n - 1)
    gf = slide.shapes.add_table(n, m, x, y, w, h)
    tbl = gf.table
    tblPr = gf._element.graphic.graphicData.tbl.tblPr
    tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = int(w * cw / total)
    for i in range(n):
        tbl.rows[i].height = Inches(hdr_h if i == 0 else row_h)
        for j in range(m):
            cell = tbl.cell(i, j)
            v = rows[i][j]
            st = {}
            if isinstance(v, tuple):
                v, st = v
            cell.margin_left = cell.margin_right = Inches(0.05)
            cell.margin_top = cell.margin_bottom = Inches(0.01)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            fill = st.get("fill")
            if fill is None and row_fill:
                fill = row_fill(i)
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill if fill is not None else (WHITE if i else CARD)
            tf = cell.text_frame
            tf.word_wrap = True
            paras = v if isinstance(v, list) else [v]
            for pi, para in enumerate(paras):
                p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
                a = (align[j] if align else ("l" if j == 0 else "r"))
                p.alignment = {"l": PP_ALIGN.LEFT, "r": PP_ALIGN.RIGHT, "c": PP_ALIGN.CENTER}[a]
                runs = para if isinstance(para, list) else [(para, {})]
                for t, rs in runs:
                    r = p.add_run()
                    r.text = t
                    f = r.font
                    f.name = MONO if rs.get("mono", st.get("mono", False)) else FONT
                    f.size = Pt(rs.get("size", st.get("size", header_size if i == 0 else body_size)))
                    f.bold = rs.get("bold", st.get("bold", i == 0))
                    f.color.rgb = rs.get("color", st.get("color", hdr_color if i == 0 else TEXT))
            if cell_style:
                cell_style(i, j, cell)
    # 细边框
    from pptx.oxml.ns import qn
    from lxml import etree
    for i in range(n):
        for j in range(m):
            tcPr = tbl.cell(i, j)._tc.get_or_add_tcPr()
            for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                ln = etree.SubElement(tcPr, qn(tag), w="6350" if tag == "a:lnB" else "0")
                if tag == "a:lnB":
                    sf = etree.SubElement(ln, qn("a:solidFill"))
                    etree.SubElement(sf, qn("a:srgbClr"), val=str(bd))
                else:
                    etree.SubElement(ln, qn("a:noFill"))
    return gf, h


def header(slide, num, title_runs, sub, size=16):
    # 顶部渐变条
    bar = rect(slide, 0, 0, SLIDE_W, Inches(0.06), fill=ACCENT, line=None, shape=MSO_SHAPE.RECTANGLE)
    bar.fill.gradient(); bar.fill.gradient_angle = 0
    st = bar.fill.gradient_stops
    st[0].color.rgb = ACCENT; st[0].position = 0
    st[1].color.rgb = GOOD; st[1].position = 1
    runs = [(f"{num}  ", {"size": 11, "color": MUTED, "bold": True})] + title_runs
    text(slide, Inches(0.35), Inches(0.16), Inches(12.6), Inches(0.42), runs, size=size, bold=True, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    text(slide, Inches(0.35), Inches(0.56), Inches(12.6), Inches(0.26), sub, size=8, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)


def footer(slide, left, right):
    text(slide, Inches(0.35), Inches(7.02), Inches(10.1), Inches(0.42), left, size=6.5, color=MUTED, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
    text(slide, Inches(10.5), Inches(7.02), Inches(2.5), Inches(0.42), right, size=6.5, color=MUTED, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


LV = {"prod": ("产品级", PROD), "sess": ("会话级", SESS), "dyn": ("每 Q 重写", BAD), "dynt": ("随 U 追加", BAD), "q": ("query 级", QLV), "loop": ("loop 级", LOOP)}


# ---------------- P1 ----------------
def slide_p1(prs, D):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    b = D["base"]
    header(s, "01", [("数据基线:", {}), ("优化前 / 优化后各模块长度", {"color": ACCENT}), ("——全稿所有数字只从这一页取", {})],
           "K token · DS-V4 tokenizer · 两边同一份内容、只换排布,同名模块长度相同;序号 = 送给模型的物理顺序;蓝底行 = 位置变了")

    T3 = f"{b['web']:g}/{b['skill']:g}/{b['tool']:g}K"
    left_rows = [
        (1, "System Prompt·固定", "prod", "14 个 ## 段:角色 / 原则 / 规范 / 安全 / TOOLS.md", K1(b["sp"]), False),
        (2, "System Prompt·动态", "dyn", "available skills 3 + 动态记忆 / 知识经验 0.3 + 前台应用 / 小艺状态 / 位置 0.05,写在 System Prompt 中部", K1(b["dyn"]), True),
        (3, "会话级上下文", "sess", "user.md 0.5 + device_info 0.5 + memory.md(长期记忆)0.5,System Prompt 尾", K1(b["sess"]), False),
        (4, "Tools", "prod", "12 个工具 schema,chat template 拼在 System Prompt 之后", K1(b["tools"]), True),
        (5, "对话历史 1–10 轮", "sess", "上一会话的对话历史,会话创建时载入、大概率不在缓存,UAT 头", K1(b["hist"]), False),
        (6, "UAT-U · query", "q", "用户原话,不含注入", K1(b["u"]), False),
        (7, "UAT-A · tool call", "loop", "函数名 + 参数;含检索 / skill preload 伪造的 tool call", K1(b["call"]), False),
        (8, "UAT-T · 工具结果", "loop", "检索(伪造 T)/ skill preload(伪造 T)/ 正常工具;伪造 A/T 随 U 进入、不花模型调用", T3, False),
        (9, "UAT-A · 最终 answer", "q", "该 Q 最后一次回复,不再调工具", K1(b["ans"]), False),
    ]
    right_rows = [
        (1, "Tools", "prod", "全量 12 个预置,排到最前", K1(b["tools"]), True),
        (2, "System Prompt·固定", "prod", "同 14 段;不塞高频 skill", K1(b["sp"]), False),
        (3, "会话级上下文", "sess", "user.md 0.5 + device_info 0.5 + memory.md(长期记忆)0.5,System Prompt 尾;会话内不变", K1(b["sess"]), False),
        (4, "对话历史 1–10 轮", "sess", "上一会话的对话历史,会话创建时载入、大概率不在缓存 → 会话首调必算一次、预热可覆盖;UAT 头", K1(b["hist"]), False),
        (5, "UAT-U · 动态", "dynt", "available skills 3 + 动态记忆 / 知识经验 0.3 + 前台应用 / 小艺状态 / 位置 0.05,当轮 user 消息前半,不回写 System Prompt", K1(b["dyn"]), True),
        (6, "UAT-U · query", "q", "当轮 user 消息后半,永远在动态之后", K1(b["u"]), False),
        (7, "UAT-A · tool call", "loop", "函数名 + 参数;含检索 / skill preload 伪造的 tool call", K1(b["call"]), False),
        (8, "UAT-T · 工具结果", "loop", "检索(伪造 T)/ skill preload(伪造 T)/ 正常工具;伪造 A/T 随 U 进入、不花模型调用", T3, False),
        (9, "UAT-A · 最终 answer", "q", "该 Q 最后一次回复,不再调工具", K1(b["ans"]), False),
    ]

    def baseline_table(x, w, title, sub, rows):
        y0 = Inches(0.9)
        card(s, x, y0, w, Inches(6.05), title, sub)
        y = y0 + Inches(0.4)
        trs = [["序号", "模块", ("长度", {})]]
        for (i, name, lv, desc, ln, mv) in rows:
            lvname, lvcol = LV[lv]
            cell = [[(name, {"bold": True, "size": 8.5, "color": ACCENT if mv else TEXT}), ("  " + lvname, {"size": 6, "bold": True, "color": lvcol})],
                    [(desc, {"size": 6.3, "color": MUTED})]]
            fill = MV_BG if mv else WHITE
            trs.append([(str(i), {"fill": fill, "bold": True, "color": ACCENT, "mono": True, "size": 8.5}),
                        (cell, {"fill": fill}),
                        (ln, {"fill": fill, "mono": True, "size": 8.5, "bold": True})])
        table(s, x + Inches(0.08), y, w - Inches(0.16), [0.62, 4.85, 1.5], trs, row_h=0.585, hdr_h=0.24, align=["c", "l", "r"])

    baseline_table(Inches(0.35), Inches(4.3), "优化前(当前方案)", "物理序 System Prompt → Tools → UAT", left_rows)
    baseline_table(Inches(4.78), Inches(4.3), "优化后(目标方案)", "物理序 Tools → System Prompt → UAT", right_rows)

    # 右栏:业务占比
    rx, rw = Inches(9.21), Inches(3.78)
    y = card(s, rx, Inches(0.9), rw, Inches(1.72), "业务占比", "按 Q 统计 · 决定每轮必算与沉淀")
    mix = b["mix"]
    biz = [["类型", "占比", "调用链"],
           [("问答", {"bold": True}), (pct(mix["qa"]), {"mono": True}), (f"U → 伪造 A/T(检索 {b['web']:g}K)→ 调① answer · 1 次", {"size": 6.5, "color": MUTED})],
           [("任务", {"bold": True}), (pct(mix["task"]), {"mono": True}), (f"U → 伪造 A/T(skill {b['skill']:g}K)→ 调① tool call → T 工具 {b['tool']:g}K → 调② answer · 2 次", {"size": 6.5, "color": MUTED})],
           [("闲聊", {"bold": True}), (pct(mix["chat"]), {"mono": True}), ("U → 调① answer · 1 次", {"size": 6.5, "color": MUTED})]]
    table(s, rx + Inches(0.08), y, rw - Inches(0.16), [0.7, 0.7, 4.2], biz, row_h=0.35, hdr_h=0.22, align=["l", "r", "l"], body_size=7.5)

    # 右栏:session 轮数分布(横向柱)
    y0 = Inches(2.72)
    card(s, rx, y0, rw, Inches(4.23), "session 轮数分布", f"实测 {round(D['sessionN'] / 1000)}k 条 · session = 拉起 app 到锁屏 / 超时重置")
    groups = []
    for g in D["p1Groups"]:
        sh = sum(r["share"] for r in D["rounds"] if r["name"] in g[1])
        cnt = sum(r["raw"] for r in D["rounds"] if r["name"] in g[1])
        nmin = min(r["n"] for r in D["rounds"] if r["name"] in g[1])
        groups.append((g[0], sh, cnt, nmin))
    import math
    mx = math.sqrt(groups[0][1])
    bx, by, bw_total = rx + Inches(0.12), y0 + Inches(0.45), rw - Inches(0.24)
    lab_w, val_w = Inches(0.62), Inches(1.05)
    track_w = bw_total - lab_w - val_w - Inches(0.12)
    rh = Inches(0.33)
    for i, (name, sh, cnt, nmin) in enumerate(groups):
        yy = by + rh * i
        text(s, bx, yy, lab_w, rh, [(name, {"bold": True, "size": 7.5})], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        col = PROD if nmin <= 5 else (SESS if nmin <= 9 else RGBColor(0x5B, 0x21, 0xB6))
        w = max(math.sqrt(sh) / mx, 0.012) * track_w
        rect(s, bx + lab_w + Inches(0.06), yy + rh * 0.28, int(w), int(rh * 0.44), fill=col, line=None, radius=0.3)
        val = pct1(sh) if sh >= 0.01 else f"{sh * 100:.2f}%"
        text(s, bx + lab_w + Inches(0.06) + track_w + Inches(0.06), yy, val_w, rh,
             [(val, {"mono": True, "size": 7.5, "bold": True}), (f"  {cnt:,}", {"mono": True, "size": 6, "color": MUTED})], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    # 均值高亮行
    ay = by + rh * len(groups) + Inches(0.1)
    ab = rect(s, bx, ay, bw_total, Inches(0.42), fill=BANNER_BG, line=BANNER_BD, radius=0.15)
    text(s, 0, 0, 0, 0, [(f"{D['sessAvg']:.2f}", {"mono": True, "bold": True, "size": 12, "color": ACCENT}), (" 轮 / session", {"size": 8, "bold": True}),
                         ("    ·    ", {"color": MUTED}),
                         (f"{D['dayAvgRounds']:.2f}", {"mono": True, "bold": True, "size": 12, "color": ACCENT}), (" 轮 / 日", {"size": 8, "bold": True}), ("(用户 × 日)", {"size": 6.5, "color": MUTED})],
         anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, shape=ab)
    text(s, bx, ay + Inches(0.48), bw_total, Inches(0.5),
         f"条长按平方根缩放(1 轮占 {pct1(groups[0][1])},线性画会压扁其余),读数看右侧数字;P3 计算第 {D['trimAt']} 轮起按上下文裁剪计、10 轮+ 按 12 轮。色系与 P3 分桶一致:蓝 1–5 轮 · 紫 6–9 轮 · 深紫 10 轮+。",
         size=6.3, color=MUTED, line_spacing=1.05)

    footer(s, "口径:Tools / System Prompt·固定 / 会话级 / 动态 / query / answer 取现网实测(token_breakdown.csv、prompt段内容变化对照);T 三类与 tool call 为工作值,两边同长。会话内缓存不被逐出、无裁剪;命中按块上报的粒度误差(DS 4K)不计。本页落档 16(优化后)/ 17(优化前)。",
           "KV Cache 方案(汇报版)· PPT v2 · P1 数据基线")


# ---------------- P2 ----------------
ROWS_CUR = ["H N F D F N N . . . . . . . .", "H R D D D D D D N N . . . . .", "H H H H H H H H H H N . . . .", "H R D D D D D D D D D D N . .", "H R D D D D D D D D D D D D N"]
ROWS_TGT = ["H H F F N N . . . . . . . .", "H H H H H H H N N . . . . .", "H H H H H H H H H N . . . .", "H H H H H H H H H H H N . .", "H H H H H H H H H H H H H N"]
CELL_COL = {"H": HIT, "R": RW, "D": DRAG, "N": NEW, "F": FIRST, ".": NA}


def slide_p2(prs, D):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    b = D["base"]; demo = D["demo"]; agg = D["agg"]
    header(s, "02", [("上下文组成对比:", {}), ("当前 vs 目标", {"color": ACCENT}), ("——Tools 移到最前成为产品级,动态下沉 U 随轮追加", {})],
           "数字全部来自 P1;会话级只在会话首调算一次,其后只算当轮新增;段宽为示意、不按比例")

    # ---- 结构条 ----
    y0 = Inches(0.9)
    card(s, Inches(0.35), y0, Inches(12.63), Inches(1.5), "结构对照:按“变化频率”分四级",
         "物理序:当前 System Prompt → Tools → UAT;目标 Tools → System Prompt → UAT")
    # 图例
    lg = [("产品级", PROD), ("会话级", SESS), ("query 级", QLV), ("loop 级", LOOP), ("动态 · 每 Q 原位重写(病灶)", BAD)]
    lx = Inches(8.3)
    for lab, col in lg:
        rect(s, lx, y0 + Inches(0.12), Inches(0.14), Inches(0.1), fill=col, line=None, radius=0.2)
        text(s, lx + Inches(0.16), y0 + Inches(0.06), Inches(1.7), Inches(0.22), lab, size=6, color=MUTED, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        lx += Inches(0.72 if len(lab) <= 4 else (0.8 if len(lab) < 8 else 1.6))

    def bar_row(y, who, sub, segs):
        text(s, Inches(0.45), y, Inches(0.7), Inches(0.5), [[(who, {"bold": True, "size": 9})], [(sub, {"size": 6, "color": MUTED})]], anchor=MSO_ANCHOR.MIDDLE)
        x = Inches(1.2); total_w = Inches(11.65); tot = sum(w for _, _, _, w, _ in segs)
        for (label, col, subt, w, outline) in segs:
            sw = int(total_w * w / tot) - Inches(0.03)
            r = rect(s, x, y, sw, Inches(0.27), fill=(WHITE if outline else col), line=(BAD if outline else None), radius=0.2, lw=1.25)
            text(s, 0, 0, 0, 0, [(label, {"size": 6.8, "bold": True, "color": (BAD if outline else WHITE)})], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=r, wrap=False)
            text(s, x, y + Inches(0.27), sw, Inches(0.2), subt, size=5.8, color=MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP, wrap=False)
            x += sw + Inches(0.03)

    bar_row(y0 + Inches(0.4), "当前", "P1 左表 · 1–9", [
        (f"① System Prompt·固定 {K1(b['sp'])}", PROD, "14 段 · 产品级 · 命中(断点实测停在这里)", 4.2, False),
        (f"② System Prompt·动态⚡ {K1(b['dyn'])}", BAD, "skills / memory · 每 Q 原位重写", 1.75, False),
        (f"③ 会话级 {K1(b['sess'])}", SESS, "user.md + device · 被连带", 1.1, False),
        (f"④ Tools {K1(b['tools'])}", PROD, "12 个工具 · 被连带", 1.15, False),
        (f"⑤ 历史 10 轮 {K1(b['hist'])}", SESS, "完整 U/A/T · 被连带", 1.3, False),
        ("⑥–⑨ 本会话 U/A/T(随轮增长)", QLV, "U 仅 query · 每 Q 首调被连带重算", 2.6, False)])
    bar_row(y0 + Inches(0.9), "目标", "P1 右表 · 1–9", [
        (f"① Tools {K1(b['tools'])}", PROD, "全量 12 个预置 · 全网命中", 1.15, False),
        (f"② System Prompt·固定 {K1(b['sp'])}", PROD, "14 段不变、不塞 skill · 命中", 4.2, False),
        (f"③ 会话级 {K1(b['sess'])}", SESS, "会话内冻结 · 首调后命中", 1.1, False),
        (f"④ 历史 10 轮 {K1(b['hist'])}", SESS, "完整 U/A/T · 首调后命中", 1.3, False),
        (f"⑤ U·动态 {K1(b['dyn'])}", BAD, "↑动态⚡ 下沉 · 随 U 追加、不回写", 1.75, True),
        ("⑥–⑨ 本会话 U/A/T(随轮增长)", QLV, "纯追加,前缀永不改写", 2.6, False)])
    text(s, Inches(0.45), y0 + Inches(1.27), Inches(12.4), Inches(0.22),
         [("目标行怎么来的 · 三步改造:", {"bold": True}), ("前缀区只放会话创建时已确定的内容——① 动态从 System Prompt 下沉到当轮 U 前半、随 U 追加(已上线);② 全量 12 个工具预置并移到最前(需 chat template);③ 动态只追加、不回写 System Prompt。不新增预置段。", {})],
         size=6.5, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)

    # ---- 逐次调用矩阵 ----
    y1 = Inches(2.47)
    card(s, Inches(0.35), y1, Inches(12.63), Inches(3.0), "同一会话逐次调用展开:行 = 一次模型调用,列 = 上下文位置,格色 = 缓存命运",
         f"Q1 问答 · Q2 任务 · Q3 / Q4 闲聊,共 {demo['nCalls']} 次调用;检索 / skill 以伪造 A/T 随 U 进入、不占调用", title_size=9.5)
    lg2 = [("已缓存", HIT), ("原位重写", RW), ("被连带·白算", DRAG), ("新增·必算", NEW), ("会话首调载入", FIRST)]
    lx = Inches(9.4)
    for lab, col in lg2:
        rect(s, lx, y1 + Inches(0.12), Inches(0.14), Inches(0.1), fill=col, line=None, radius=0.2)
        text(s, lx + Inches(0.16), y1 + Inches(0.06), Inches(1.0), Inches(0.22), lab, size=6, color=MUTED, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        lx += Inches(0.6 if len(lab) <= 4 else 0.78)

    uat_cur = ["U1", "A₁T₁ 检索", "ans₁", "U2", "A₂T₂ skill", "A₂′T₂′ 工具", "ans₂", "U3", "ans₃", "U4"]
    uat_tgt = ["Δ+U1", "A₁T₁ 检索", "ans₁", "Δ+U2", "A₂T₂ skill", "A₂′T₂′ 工具", "ans₂", "Δ+U3", "ans₃", "Δ+U4"]
    pre_cur = [("① 固定", K1(b["sp"])), ("② 动态⚡", K1(b["dyn"])), ("③ 会话级", K1(b["sess"])), ("④ Tools", K1(b["tools"])), ("⑤ 历史", K1(b["hist"]))]
    pre_tgt = [("① Tools", K1(b["tools"])), ("② 固定", K1(b["sp"])), ("③ 会话级", K1(b["sess"])), ("④ 历史", K1(b["hist"]))]

    def acct(c, side):
        st = c[side]; out = []
        if st["waste"] > 0:
            p = st.get("parts") or {}
            lab = (f"{K1(p['tools'] + p['sess'] + p['hist'])[:-1]} + 本会话史 {K1(p['own'])[:-1]}" if p.get("own") else "Tools")
            out += [(f"白算 {K1(st['waste'])[:-1]}", {"color": WARN, "bold": True}), (f"({lab}) + ", {"color": MUTED, "size": 5.8})]
        elif side == "tgt" and c["kind"] != "loop":
            out += [("白算 0", {"color": GOOD, "bold": True}), (" + ", {})]
        if st["first"] > 0:
            out += [(f"首调载入 {K1(st['first'])[:-1]}", {"color": ACCENT, "bold": True}), (" + ", {})]
        out += [(f"必算 {K1(st['nw'])[:-1]}", {"color": ACCENT, "bold": True}),
                (f"  = {K1(st['real'])[:-1]}K", {"bold": True, "mono": True, "color": (BAD if side == "cur" and st["waste"] > 0 else GOOD)}),
                (f"   命中 {pct(st['hit'] / st['input'])}", {"color": TEXT, "size": 6, "bold": True})]
        return out

    def matrix(y, title, pre, uat, rows, side, hit_all):
        cols = len(pre) + len(uat)
        x0 = Inches(1.35); grid_w = Inches(7.6); acc_x = x0 + grid_w + Inches(0.1); acc_w = Inches(3.5)
        cw = grid_w / cols
        text(s, Inches(0.45), y, Inches(0.88), Inches(0.3), [[(title, {"bold": True, "size": 8.5})], [(f"{demo['nCalls']} 次加权命中 {pct(hit_all)}", {"size": 5.4, "color": MUTED})]], anchor=MSO_ANCHOR.MIDDLE)
        # 列头
        for j in range(cols):
            if j < len(pre):
                lab = [[(pre[j][0], {"bold": True, "size": 5.6})], [(pre[j][1], {"size": 5.6, "color": (BAD if "动态" in pre[j][0] else MUTED), "mono": True})]]
            else:
                lab = [[(uat[j - len(pre)], {"bold": True, "size": 5.6})]]
            text(s, x0 + int(cw * j), y, int(cw), Inches(0.3), lab, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.0, wrap=True)
        text(s, acc_x, y, acc_w, Inches(0.3), "本次调用账目:白算 + 必算 = 实算(K)· 命中率", size=5.8, color=MUTED, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        # 分组标签线
        rh = Inches(0.19)
        for i, line in enumerate(rows):
            yy = y + Inches(0.31) + rh * i
            c = demo["calls"][i]
            text(s, Inches(0.45), yy, Inches(0.88), rh, [[(c["t1"], {"bold": True, "size": 6.3})], [(c["t2"], {"size": 5.2, "color": MUTED})]], anchor=MSO_ANCHOR.MIDDLE, margin=0.0)
            for j, k in enumerate(line.split(" ")):
                rect(s, x0 + int(cw * j) + Inches(0.012), yy + Inches(0.03), int(cw) - Inches(0.024), rh - Inches(0.06), fill=CELL_COL[k], line=None, radius=0.15)
            text(s, acc_x, yy, acc_w, rh, acct(c, side), size=6.3, anchor=MSO_ANCHOR.MIDDLE, margin=0.0, wrap=False)
        # UAT 分隔线
        ln = s.shapes.add_connector(1, x0 + int(cw * len(pre)), y + Inches(0.3), x0 + int(cw * len(pre)), y + Inches(0.31) + rh * len(rows))
        ln.line.color.rgb = QLV; ln.line.width = Pt(1)

    matrix(y1 + Inches(0.4), "当前", pre_cur, uat_cur, ROWS_CUR, "cur", demo["curHit"])
    matrix(y1 + Inches(1.7), "目标", pre_tgt, uat_tgt, ROWS_TGT, "tgt", demo["tgtHit"])

    # ---- 账目逻辑三格 ----
    y2 = Inches(5.54)
    eqs = [
        ([("① 输入 = ", {}), ("已缓存", {"color": GOOD}), (" + ", {}), ("白算", {"color": BAD}), (" + ", {}), ("必算", {"color": ACCENT})],
         "已缓存 = 前缀逐 token 一致的深度;白算 = 本该命中却因前缀改写而重算;必算 = 本次必要新增"),
        ([("② ", {}), ("实算", {"color": WARN}), ("(内部口径“硬算”)= ", {}), ("白算", {"color": BAD}), (" + ", {}), ("必算", {"color": ACCENT}), (" = 输入 − 已缓存", {})],
         "prefill 真正计算的量,引擎逐请求实报;账单与 TTFT 都跟它走,成本收益看它降多少"),
        ([("③ cache 命中率 = ", {}), ("已缓存", {"color": GOOD}), (" ÷ 输入", {})],
         "业界通用健康度;分母含必算,必算越多命中率越低,不如实算直观,按业界口径一并报"),
    ]
    ew = Inches(4.13)
    for i, (runs, note) in enumerate(eqs):
        x = Inches(0.35) + (ew + Inches(0.12)) * i
        rect(s, x, y2, ew, Inches(0.5), fill=WHITE)
        text(s, x + Inches(0.08), y2 + Inches(0.02), ew - Inches(0.16), Inches(0.22), runs, size=8, bold=True, mono=True, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        text(s, x + Inches(0.08), y2 + Inches(0.24), ew - Inches(0.16), Inches(0.24), note, size=5.8, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)

    # ---- 四把尺子 ----
    y3 = Inches(6.1)
    fn = D["firstNew"]
    s4 = D["sessN"][3]
    rulers = [
        ([("目标 ① ", {}), ("白算 → 0", {"color": BAD})], "验收线 · 排布改造唯一可控项",
         [(K1(demo["curWaste"]), {"color": BAD}), (" → ", {"color": MUTED, "size": 8}), ("0", {"color": GOOD}), (f"  上表 {demo['nCalls']} 次调用合计", {"size": 6, "color": GOOD})],
         "Tools / 会话级 / 历史 / 本会话已写被前缀改写连带;实测 Q≥2 首调 miss 21–34K、白算占 90–99%。目标态 loop 白算 > 1 块即报警"),
        ([("目标 ② ", {}), ("必算尽可能小", {"color": ACCENT})], "整轮必要新增 · 本次不记成绩 · 下一步",
         [(K1(demo["tgtNew"]), {"color": ACCENT}), (" = ", {"color": MUTED, "size": 8}), (K1(demo["tgtNew"]), {"color": ACCENT}), ("  两边相同", {"size": 6, "color": ACCENT})],
         f"首调必算 闲聊 {K1(fn['chat'])} / 任务 {K1(fn['task'])}(含 skill)/ 问答 {K1(fn['qa'])}(含检索),决定 TTFT;整轮实算 {K1(D['Tn'])} 决定成本。下一步:available skills 升会话级、检索 / 工具结果精简"),
        ([("成本看 ", {}), ("实算", {"color": WARN}), (" 降幅", {})], "本页示例 4 轮 session · 白算归零后实算 = 必算",
         [(K0(demo["curTot"]), {"color": BAD}), (" → ", {"color": MUTED, "size": 8}), (K0(demo["tgtTot"]), {"color": GOOD}), (f"  −{pct(demo['saveRate'])}", {"size": 7, "color": GOOD})],
         f"省额 {K1(demo['curTot'] - demo['tgtTot'])} = 白算;Q≥2 首调 {K1(demo['qFirstCurLo'])}–{K1(demo['qFirstCurHi'])} → {K1(demo['qFirstTgtLo'])}–{K1(demo['qFirstTgtHi'])},与轮数脱钩。业务加权 4 轮 session −{pct(s4['save'])}、全业务(均 {D['sessAvg']:.2f} 轮)−{pct(agg['saveRate'])},见 P3"),
        ([("cache 命中率", {})], "已缓存 ÷ 输入 · 全业务加权",
         [(pct(agg["curHitAll"]), {"color": BAD}), (" → ", {"color": MUTED, "size": 8}), (pct(agg["tgtHitAll"]), {"color": GOOD}), ("  业务 × session 分布加权", {"size": 6, "color": GOOD})],
         f"当前断点恒在 System Prompt·固定末(实测约 60%,与推演相符)。目标受必算拖累——会话首调 {K1(b['sess'] + b['hist'])}、检索 {K1(b['web'])}、skill {K1(b['skill'])} 都在分母:6 轮 {pct(D['hit']['h6'])}、12 轮 {pct(D['hit']['h12'])},1Q 只 {pct(D['hit']['h1'])},而 {pct(agg['share1'])} 的 session 只有 1 轮"),
    ]
    rw = Inches(3.07)
    for i, (t, sub, val, note) in enumerate(rulers):
        x = Inches(0.35) + (rw + Inches(0.12)) * i
        rect(s, x, y3, rw, Inches(0.87))
        text(s, x + Inches(0.08), y3 + Inches(0.02), rw - Inches(0.16), Inches(0.2), t + [("  " + sub, {"size": 5.4, "color": MUTED, "bold": False})], size=7.5, bold=True, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        text(s, x + Inches(0.08), y3 + Inches(0.2), rw - Inches(0.16), Inches(0.22), val, size=10.5, bold=True, mono=True, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        text(s, x + Inches(0.08), y3 + Inches(0.42), rw - Inches(0.16), Inches(0.44), note, size=5.4, color=MUTED, line_spacing=1.0)

    footer(s, "数字全部来自 P1 基线,逐调用推演;当前侧与三条实测会话一致(断点恒在 System Prompt·固定末、Q 首调 miss 14–34K、loop 白算 ≤1 块)。目标侧输入比当前略长(动态留在历史里),是只追加的代价,已计入命中率分母。目标序 Tools → System Prompt 需 chat template 支持(见 06)。",
           "KV Cache 方案(汇报版)· PPT v2 · P2")


# ---------------- P3 ----------------
def slide_p3(prs, D):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    b = D["base"]; agg = D["agg"]
    header(s, "03", [("轮数 × 业务:", {}), (f"当前每深一轮多烧 {K1(D['alpha'])},目标 Q2 起钉在 {K1(D['Tn'])}", {"color": ACCENT}),
                     (f"——全量加权 token −{pct(agg['saveRate'])}(成本 −{pct(agg['costSave'])}),cache 命中率 {pct(agg['curHitAll'])} → {pct(agg['tgtHitAll'])}", {})],
           f"左 = 第 n 轮(单轮)· 右 = L 轮的 session 与全量;业务 47 / 28 / 25 加权,第 {D['trimAt']} 轮起上下文裁剪", size=13.5)

    colw = Inches(6.25); gap = Inches(0.13)
    lx = Inches(0.35); rx = lx + colw + gap
    top = Inches(0.9); card_h = Inches(5.05)
    key_y = top + Inches(0.38); key_h = Inches(0.62)
    chart_y = key_y + key_h + Inches(0.06); chart_h = Inches(2.2)
    tblh_y = chart_y + chart_h + Inches(0.02)
    tbl_y = tblh_y + Inches(0.24)

    def dim_title(x, tag, title, sub):
        chip(s, x + Inches(0.1), top + Inches(0.09), Inches(1.15), Inches(0.2), tag, RGBColor(0xDB, 0xE4, 0xFB), color=ACCENT, size=7)
        text(s, x + Inches(1.3), top + Inches(0.04), colw - Inches(1.4), Inches(0.3), [(title, {"bold": True, "size": 10.5}), ("  " + sub, {"size": 7, "color": MUTED})], anchor=MSO_ANCHOR.MIDDLE, wrap=False)

    # ================= 左列:第 n 轮 =================
    rect(s, lx, top, colw, card_h)
    dim_title(lx, "第 n 轮", "单轮视角", "一轮 = 一个 Q,含该轮全部调用")
    kb = rect(s, lx + Inches(0.1), key_y, colw - Inches(0.2), key_h, fill=WHITE)
    fn = D["firstNew"]; fe = D["firstExtra"]; qn = D["qNew"]
    text(s, 0, 0, 0, 0, [
        [("目标每轮实算(Q≥2)= 首调必算 + 第二次调用:", {"bold": True, "color": TEXT}), ("闲聊 ", {}), (K1(fn["chat"]), {"bold": True, "color": GOOD}),
         (f"(动态 {K1(b['dyn'])} + query {K1(b['u'])} + 上轮 answer {K1(b['ans'])})| 任务 {K1(fn['chat'])} + skill {K1(fe['task'])} + 工具 {K1(D['loopNew']['task'])} = ", {}), (K1(qn["task"]), {"bold": True, "color": GOOD}),
         (f" | 问答 {K1(fn['chat'])} + 检索 {K1(fe['qa'])} = ", {}), (K1(qn["qa"]), {"bold": True, "color": GOOD}), (" → 按 47 / 28 / 25 加权 = ", {}), (K1(D["Tn"]), {"bold": True, "color": GOOD, "size": 9, "mono": True}), (",与轮数无关。", {})],
        [(f"当前每轮 = {K1(D['Tn'])} + 白算:", {"bold": True, "color": TEXT}), ("固定项 ", {}), (K1(D["fixed2"]), {"bold": True, "color": BAD}), ("(Tools + 会话级 + 历史)+ 本会话史 ", {}), (K1(D["alpha"]), {"bold": True, "color": BAD}),
         (f" × (n−1),每深一轮多烧 {K1(D['alpha'])};第 {D['trimAt']} 轮起裁剪封顶。", {})],
    ], size=6.6, color=MUTED, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.1, shape=kb)

    cd = CategoryChartData()
    pr = D["perRound"]
    cd.categories = [f"第 {p['n']} 轮  −{pct(p['eta'])}" for p in pr]
    cd.add_series("当前每轮实算 C(n)", [round(p["C"], 1) for p in pr])
    cd.add_series("目标每轮实算 T(n)", [round(p["T"], 1) for p in pr])
    gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, lx + Inches(0.1), chart_y, colw - Inches(0.2), chart_h, cd)
    ch = gf.chart
    ch.font.size = Pt(7); ch.font.name = FONT
    ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.TOP; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(7)
    ch.value_axis.has_major_gridlines = True; ch.value_axis.major_gridlines.format.line.color.rgb = GRID
    ch.value_axis.tick_labels.font.size = Pt(7); ch.value_axis.tick_labels.number_format = '0"K"'; ch.value_axis.tick_labels.number_format_is_linked = False
    ch.value_axis.maximum_scale = 60; ch.value_axis.minimum_scale = 0; ch.value_axis.format.line.fill.background()
    ch.category_axis.tick_labels.font.size = Pt(7.5); ch.category_axis.tick_labels.font.bold = True; ch.category_axis.format.line.color.rgb = GRID
    plot = ch.plots[0]; plot.has_data_labels = True
    plot.data_labels.font.size = Pt(7); plot.data_labels.font.bold = True; plot.data_labels.number_format = '0.0'; plot.data_labels.number_format_is_linked = False
    for ser, col, pos in zip(plot.series, (BAD, GOOD), (XL_LABEL_POSITION.ABOVE, XL_LABEL_POSITION.BELOW)):
        ser.format.line.color.rgb = col; ser.format.line.width = Pt(2.25); ser.smooth = False
        ser.marker.style = XL_MARKER_STYLE.CIRCLE; ser.marker.size = 6
        ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = col; ser.marker.format.line.color.rgb = col
        ser.data_labels.font.color.rgb = col; ser.data_labels.position = pos; ser.data_labels.font.size = Pt(7); ser.data_labels.show_value = True

    text(s, lx + Inches(0.1), tblh_y, colw - Inches(0.2), Inches(0.22), [("逐轮账", {"bold": True, "size": 8.5}), (f"  TTFT = 首调实算 × ≈{round(D['msPerK'])} ms/K(参考实测 14K ≈ 640 ms);预热后统一 {D['preheatMs']} ms;问答 / 任务另 +检索 / skill 两边相同", {"size": 6, "color": MUTED})], anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    rows = [["第 n 轮", "当前实算", "目标实算", "节省率 η", "该轮命中率 当前 → 目标", "首调 TTFT 当前 → 目标", "预热后"]]
    for o in D["rounds7"]:
        tag = "  会话首调" if o["n"] == 1 else ("  裁剪" if o["n"] >= D["trimAt"] else "")
        rows.append([([[(f"第 {o['n']} 轮", {"bold": True}), (tag, {"size": 5.8, "color": MUTED})]], {}), (K1(o["C"]), {"mono": True, "color": BAD}), (K1(o["T"]), {"mono": True, "color": GOOD}),
                     (f"−{pct(o['eta'])}", {"mono": True, "color": GOOD, "bold": True}),
                     ([[(pct(o["hitCur"]), {"color": BAD, "mono": True}), (" → ", {"color": MUTED}), (pct(o["hitTgt"]), {"color": GOOD, "mono": True, "bold": True})]], {}),
                     ([[(o["fCurS"], {"color": BAD, "mono": True}), (" → ", {"color": MUTED}), (o["fTgtS"], {"color": GOOD, "mono": True})]], {}),
                     (f"{D['preheatMs']} ms", {"mono": True, "color": GOOD, "bold": True})])
    table(s, lx + Inches(0.1), tbl_y, colw - Inches(0.2), [1.25, 0.9, 0.9, 0.85, 1.5, 1.5, 0.8], rows, row_h=0.2, hdr_h=0.22, body_size=6.8, header_size=6.6,
          align=["l", "r", "r", "r", "r", "r", "r"])

    # ================= 右列:L 轮的 session =================
    rect(s, rx, top, colw, card_h)
    dim_title(rx, "L 轮的 session", "session 视角", f"一条 L 轮的 session 全部调用累加;实测 {round(D['sessionN'] / 1000)}k 条")
    kb2 = rect(s, rx + Inches(0.1), key_y, colw - Inches(0.2), key_h, fill=BANNER_BG, line=BANNER_BD)
    text(s, 0, 0, 0, 0, [
        [("全量  ", {"bold": True, "color": ACCENT, "size": 7}), (f"按实测 session 分布(均 {D['sessAvg']:.2f} 轮,{pct(agg['share1'])} 只有 1 轮)加权 → ", {"bold": True, "color": TEXT}),
         ("token ", {"color": TEXT}), (f"−{pct(agg['saveRate'])}", {"bold": True, "color": GOOD, "size": 10, "mono": True}), ("(成本 ", {"color": TEXT}), (f"−{pct(agg['costSave'])}", {"bold": True, "color": GOOD, "mono": True}),
         (")· cache 命中率 ", {"color": TEXT}), (pct(agg["curHitAll"]), {"bold": True, "color": BAD, "mono": True}), (" → ", {"color": MUTED}), (pct(agg["tgtHitAll"]), {"bold": True, "color": GOOD, "size": 10, "mono": True})],
        [("每 session 的数已按业务加权,全量 = 再按 session 长度分布加权;新架构上线后 session 变长、收益向长 session 靠。", {"size": 6.2})],
    ], size=7, color=MUTED, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.1, shape=kb2)

    sn = D["sessN"]
    cd2 = CategoryChartData()
    cd2.categories = [f"{o['n']} 轮" + ("·裁剪" if o["n"] >= D["trimAt"] else "") for o in sn]
    cd2.add_series("当前实算 / session", [round(o["cur"]) for o in sn])
    cd2.add_series("目标实算 / session", [round(o["tgt"]) for o in sn])
    gf2 = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, rx + Inches(0.1), chart_y, colw - Inches(0.2), chart_h, cd2)
    c2 = gf2.chart
    c2.font.size = Pt(7); c2.font.name = FONT
    c2.has_legend = True; c2.legend.position = XL_LEGEND_POSITION.TOP; c2.legend.include_in_layout = False; c2.legend.font.size = Pt(7)
    c2.value_axis.has_major_gridlines = True; c2.value_axis.major_gridlines.format.line.color.rgb = GRID
    c2.value_axis.tick_labels.font.size = Pt(7); c2.value_axis.tick_labels.number_format = '0"K"'; c2.value_axis.tick_labels.number_format_is_linked = False
    c2.value_axis.maximum_scale = 600; c2.value_axis.format.line.fill.background()
    c2.category_axis.tick_labels.font.size = Pt(7.5); c2.category_axis.tick_labels.font.bold = True; c2.category_axis.format.line.color.rgb = GRID
    p2 = c2.plots[0]; p2.gap_width = 70; p2.overlap = -10
    p2.has_data_labels = True; p2.data_labels.font.size = Pt(6.5); p2.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    for ser, col in zip(p2.series, (BAD, GOOD)):
        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = col; ser.format.line.fill.background()
    for i, o in enumerate(sn):
        dl = p2.series[0].points[i].data_label; dl.has_text_frame = True; dl.text_frame.text = f"−{pct(o['save'])}"
        r0 = dl.text_frame.paragraphs[0].runs[0]; r0.font.size = Pt(8); r0.font.bold = True; r0.font.color.rgb = RGBColor(0x04, 0x78, 0x57)
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        d2 = p2.series[1].points[i].data_label; d2.has_text_frame = True; d2.text_frame.text = f"{round(o['tgt'])}K"
        r1 = d2.text_frame.paragraphs[0].runs[0]; r1.font.size = Pt(6.5); r1.font.color.rgb = GOOD
        d2.position = XL_LABEL_POSITION.OUTSIDE_END

    text(s, rx + Inches(0.1), tblh_y, colw - Inches(0.2), Inches(0.22), [("按 session 长度分桶", {"bold": True, "size": 8.5}), ("  6–9 / 10+ 为组内加权,10 轮+ 按 12 轮;命中率 = session 累计;柱顶 = 降幅", {"size": 6, "color": MUTED})], anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    rows = [["L 轮的 session", "占比", "当前 / session", "目标 / session", "降幅", "占当前算力", "session 累计命中率"]]
    for g in agg["groups"]:
        m = "均 " if g["multi"] else ""
        rows.append([(g["name"], {"bold": True}), (pct1(g["share"]), {"mono": True}), (m + K1(g["cur"]), {"mono": True}), (m + K1(g["tgt"]), {"mono": True}),
                     (f"−{pct(g['save'])}", {"mono": True, "color": GOOD, "bold": True}), (pct(g["shareCost"]), {"mono": True}),
                     ([[(pct(g["hitCur"]), {"color": BAD, "mono": True}), (" → ", {"color": MUTED}), (pct(g["hitTgt"]), {"color": GOOD, "mono": True, "bold": True})]], {})])
    table(s, rx + Inches(0.1), tbl_y, colw - Inches(0.2), [1.9, 0.75, 1.1, 1.1, 0.75, 0.95, 1.45], rows, row_h=0.2, hdr_h=0.22, body_size=6.8, header_size=6.6,
          align=["l", "r", "r", "r", "r", "r", "r"])

    # ================= 底部结论四格 =================
    cy = top + card_h + Inches(0.1); chh = Inches(0.86)
    items = [
        ("bars", [("2 轮以上的 session 只占 ", {}), (pct(agg["longShare"]), {"bold": True}), (",却吃掉 ", {}), (pct(agg["longShareCost"]), {"bold": True}), (" 的算力", {"bold": True}),
                  (f",贡献 {pct(agg['longShareSave'])} 的节省;越长省得越多。", {})]),
        (K1(D["qFirstNew"]), [("会话级可提前算,首调时延与轮数脱钩。", {"bold": True}), (f"{pct(agg['share1'])} 的 session 只有 1 轮,预热后 TTFT ≈ {D['preheatMs']} ms + 业务必带。", {})]),
        ("白算 → 0", [("白算归零是验收线,必算是下一步", {"bold": True}), (f":历史 {K1(b['hist'])} 整段载入改摘要、available skills 3K 升会话级。", {})]),
        ("效果不变", [("动态下沉 U 已上线,对话效果影响不大", {"bold": True}), (";后两步只改顺序不改内容,不引入效果风险。", {})]),
    ]
    cw4 = (Inches(12.63) - Inches(0.36)) / 4
    for i, (k, runs) in enumerate(items):
        x = Inches(0.35) + (cw4 + Inches(0.12)) * i
        rect(s, x, cy, cw4, chh)
        kx_w = Inches(1.15)
        if k == "bars":
            for j, (lab, v, col) in enumerate((("会话", agg["longShare"], ACCENT), ("算力", agg["longShareCost"], BAD))):
                by = cy + Inches(0.22) + Inches(0.2) * j
                text(s, x + Inches(0.1), by, Inches(0.32), Inches(0.16), lab, size=6.5, bold=True, anchor=MSO_ANCHOR.MIDDLE, margin=0.0)
                rect(s, x + Inches(0.42), by + Inches(0.03), int(Inches(0.5) * v), Inches(0.1), fill=col, line=None, radius=0.3)
                text(s, x + Inches(0.78), by, Inches(0.4), Inches(0.16), pct(v), size=7.5, bold=True, mono=True, color=col, anchor=MSO_ANCHOR.MIDDLE, margin=0.0)
        else:
            sub = "Q≥2 首调必算" if k == K1(D["qFirstNew"]) else (f"命中率 {pct(agg['curHitAll'])} → {pct(agg['tgtHitAll'])}" if k == "白算 → 0" else "")
            text(s, x + Inches(0.1), cy, kx_w, chh, [[(k, {"bold": True, "size": 11, "color": ACCENT, "mono": True})]] + ([[(sub, {"size": 6, "color": MUTED, "bold": True})]] if sub else []), anchor=MSO_ANCHOR.MIDDLE, margin=0.0)
        text(s, x + Inches(0.1) + kx_w, cy, cw4 - Inches(0.2) - kx_w, chh, runs, size=6.8, color=MUTED, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.06)

    footer(s, f"成本按计价折算 ≈{D['costFactor']:.2f} × token 降幅。裁剪:第 {D['trimAt']} 轮起上下文裁剪,本会话史累积封顶 {D['trimAt'] - 1} 轮;10 轮+ 桶按 12 轮计。前提:会话内缓存不被逐出(实测 2/21 次 loop 全 miss 属引擎侧逐出 / 路由,需会话粘性 + TTL,单独报警)。实测对照:当前命中率约 60%(推演 {pct(agg['curHitAll'])})、平均实算约 12K/调用,目标 7K;TTFT 待 prefill 吞吐实测(E6)替换;上线后由日报(引擎 cached_tokens + 应用边界埋点)替换推演。",
           "KV Cache 方案(汇报版)· PPT v2 · P3 轮数收益")


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else None
    if not src:
        cands = glob.glob(os.path.join(HERE, "kv cache方案（汇报版）.v*.slides.html"))
        def ver(p):
            m = re.search(r"\.v([\d.]+)\.slides", p); return tuple(int(x) for x in m.group(1).split("."))
        src = sorted(cands, key=ver)[-1]
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, "kv cache方案（汇报版）.v2.pptx")
    D = load_data(src)
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    slide_p1(prs, D); slide_p2(prs, D); slide_p3(prs, D)
    prs.save(out)
    print(f"源:{os.path.basename(src)}\n输出:{out}\n数字:token −{pct(D['agg']['saveRate'])} / 成本 −{pct(D['agg']['costSave'])} / 命中 {pct(D['agg']['curHitAll'])} → {pct(D['agg']['tgtHitAll'])} / session 均 {D['sessAvg']:.2f} 轮")


if __name__ == "__main__":
    main()
